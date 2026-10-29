# -*- coding: utf-8 -*-
# update_supp_tables_v2c.py — (2026-10-22 v2c) 보충표 docx(생성기 없는 정본)의 사본에서 표 칸 숫자만 v2c 재실행 산출로 바꿈.
#  입력: 원본 = _pre_v2_rerun_261022/JKMS/JKMS_FBD/JKMS_EHEC/Supplementary/Supplementary_Tables_EHEC.docx(재실행 전 정본 사본)
#  출력: _v2_rerun_261022/JKMS/JKMS_FBD_v2c/JKMS_EHEC/Supplementary/Supplementary_Tables_EHEC.docx (기존 투고 패키지는 건드리지 않음)
#  각주·설명 문장은 고치지 않음(다음 단계 문장 수정 몫). 바뀐 칸 기록 → _v2_rerun_261022/JKMS/supp_table_changes_v2c.csv
#  다중우주(S12·S13·S14 하단·S15)는 씨드 12개가 다 끝나 PREREG_MV_summary.csv가 새로 만들어졌을 때만 바꿈.
import os, copy, glob, re, sys
import pandas as pd, numpy as np
from docx import Document
from docx.oxml.ns import qn

ROOT = os.environ.get("EHEC_ROOT",".")
RS = os.path.join(ROOT, "results_ST"); SH = os.path.join(ROOT, "EHEC", "series"); RR = os.path.join(ROOT, "_v2_rerun_261022", "JKMS")
SRC = os.path.join(ROOT, "_pre_v2_rerun_261022/JKMS/JKMS_FBD/JKMS_EHEC/Supplementary/Supplementary_Tables_EHEC.docx")
DST = os.path.join(RR, "JKMS_FBD_v2c/JKMS_EHEC/Supplementary/Supplementary_Tables_EHEC.docx")
d = Document(SRC); T = d.tables; LOG = []

def _runs(c): return [r for p in c.paragraphs for r in p.runs]
def setc(t, i, j, text, sup=None, tag=""):
    c = t.rows[i].cells[j]; old = c.text; rs = [r for r in _runs(c) if r.text]
    p = c.paragraphs[0]
    base = rs[0] if rs else (p.runs[0] if p.runs else p.add_run(""))
    rpr = copy.deepcopy(base._r.find(qn("w:rPr")))
    for pp in c.paragraphs:
        for r in list(pp.runs): r._r.getparent().remove(r._r)
    for pp in c.paragraphs[1:]: pp._p.getparent().remove(pp._p)
    r = p.add_run(text)
    if rpr is not None:
        rp = copy.deepcopy(rpr); va = rp.find(qn("w:vertAlign"))
        if va is not None: rp.remove(va)
        r._r.insert(0, rp)
    if sup:
        s = p.add_run(sup)
        rp = copy.deepcopy(rpr) if rpr is not None else None
        if rp is not None: s._r.insert(0, rp)
        s.font.superscript = True
    new = text + (sup or "")
    if old != new: LOG.append(dict(table=tag, row=i, col=j, before=old, after=new))
def addrow(t, like):
    tr = copy.deepcopy(t.rows[like]._tr); t._tbl.append(tr); return len(t.rows) - 1
def delrow(t, i): t._tbl.remove(t.rows[i]._tr)
def setrows(t, first, n_needed, tag):
    """데이터 행 수를 n_needed로 맞춤(first부터 끝까지가 데이터 행)."""
    cur = len(t.rows) - first
    while cur < n_needed: addrow(t, len(t.rows) - 1); cur += 1; LOG.append(dict(table=tag, row=len(t.rows)-1, col="", before="(행 추가)", after=""))
    while cur > n_needed: delrow(t, len(t.rows) - 1); cur -= 1; LOG.append(dict(table=tag, row=first+cur, col="", before="(행 삭제)", after=""))
f3 = lambda x: f"{x:.3f}"; ci = lambda lo, hi: f"{lo:.3f}-{hi:.3f}"; aa = lambda c: "a" if str(c).upper() == "TRUE" else None
def tf(x): return str(x).upper() == "TRUE"

T2 = pd.read_csv(os.path.join(RS, "MASTER_final_table2.csv")); T2c = T2.set_index("cov")
# 보충표의 행 이름(원고 9/3 라벨) → 변수 키. 540 라벨이 바뀌어도(예: kchs_pc1 'KCHS survey composite (PC1)') 키로 찾음
OLDLAB = pd.read_csv(os.path.join(ROOT, "_v2_promote_261021", "17_KEY_기준_MASTER_final_table2_260903.csv"))
L2C = {**dict(zip(OLDLAB.label, OLDLAB["cov"])), **dict(zip(T2.label, T2["cov"]))}
def t2row(lab): return T2c.loc[L2C[lab.replace("1,000", "1000")]]
FIT = dict(l.strip().split("=") for l in open(os.path.join(RS, "final_fit.txt")))

# S7 (표 6): EHSA 분류
E = pd.read_csv(os.path.join(RS, "EHSA_agestd_sum_esri.csv")).set_index("cat").n
cat = pd.read_csv(os.path.join(RS, "EHSA_agestd_cat_esri.csv"))
order = [("New hot spot", "New hot"), ("Sporadic hot spot", "Sporadic hot"), ("New cold spot", "New cold"), ("Sporadic cold spot", "Sporadic cold"), ("No pattern detected", "No pattern")]
t = T[6]; keep_last = [c.text for c in t.rows[-1].cells]
rows = [(lab, int(E.get(k, 0))) for lab, k in order if int(E.get(k, 0)) > 0]
setrows(t, 1, len(rows) + 1, "S7")
for i, (lab, n) in enumerate(rows, 1): setc(t, i, 0, lab, tag="S7"); setc(t, i, 1, str(n), tag="S7")
n90 = int((cat.n_hot_years >= 13).sum()); setc(t, len(rows) + 1, 0, keep_last[0], tag="S7"); setc(t, len(rows) + 1, 1, str(n90), tag="S7")
S7_MAXHOT = int(cat.n_hot_years.max()); S7_NMAX = int((cat.n_hot_years == S7_MAXHOT).sum())

# S8 (표 7): 시기별 순수 공간 모형
fs = pd.read_csv(os.path.join(RS, "MASTER_final_covid_fitstat.csv")); rr = pd.read_csv(os.path.join(RS, "MASTER_final_covid_regionRR.csv"))
hi = {"2011-2019": int(rr.hi_pre.sum()), "2020-2024": int(rr.hi_post.sum())}
t = T[7]
for i, r in enumerate(fs.itertuples(), 1):
    setc(t, i, 1, f"{int(r.N):,}", tag="S8"); setc(t, i, 2, f"{r.DIC:.1f}", tag="S8"); setc(t, i, 3, f"{r.phi:.3f}", tag="S8"); setc(t, i, 4, str(hi[r.period]), tag="S8")
S8_CORR = round(float(np.corrcoef(rr.RR_pre, rr.RR_post)[0, 1]), 2); S8_BOTH = int(((rr.hi_pre == 1) & (rr.hi_post == 1)).sum())

# S9 (표 8): COVID 상호작용
ci9 = pd.read_csv(os.path.join(RS, "MASTER_final_covid_interaction.csv")).set_index("term")
m9 = {"Beef-cattle farm density": "beef_cattle_farms", "Fiscal independence": "fiscal_independence", "Net migration": "net_migration", "Current smoking": "current_smoking",
      "Post-COVID-19 period (level shift)": "post", "Beef-cattle farm density × post": "beef_cattle_farms_x_post", "Fiscal independence × post": "fiscal_independence_x_post",
      "Net migration × post": "net_migration_x_post", "Current smoking × post": "current_smoking_x_post"}
t = T[8]
for i in range(1, len(t.rows)):
    k = m9[t.rows[i].cells[0].text.strip()]; r = ci9.loc[k]
    setc(t, i, 1, f3(r.IRR), aa(r.cred), tag="S9"); setc(t, i, 2, ci(r.lo, r.hi), tag="S9")

# S11 (표 10): 최종 20변수
t = T[10]
for i in range(1, len(t.rows)):
    lab = t.rows[i].cells[0].text.strip()
    if not lab or lab[0].isdigit(): continue
    r = t2row(lab)
    setc(t, i, 2, f3(r.IRR), aa(r.cred), tag="S11"); setc(t, i, 3, ci(r.lo, r.hi), tag="S11"); setc(t, i, 4, f3(r.P_gt1), tag="S11")

# S14 (표 13) 위쪽: 음성대조 개별 추가
nc = pd.read_csv(os.path.join(RS, "MASTER_final_negcontrol.csv")).set_index("cov")
t = T[13]
for i, k in [(2, "swine_density"), (3, "poultry_density"), (4, "walking_practice")]:
    r = nc.loc[k]; setc(t, i, 2, f"{f3(r.IRR)} ({ci(r.lo, r.hi)})", tag="S14")

# S16 (표 15): 한육우 노출원
cb = pd.read_csv(os.path.join(RS, "MASTER_census_beef_check.csv")); mb = pd.read_csv(os.path.join(RS, "MASTER_monitor_beef_sens.csv"))
def gcb(model): r = cb[(cb.model == model) & cb["cov"].str.contains("beef")].iloc[0]; return r
def gmb(model, cov): return mb[(mb.model == model) & (mb["cov"] == cov)].iloc[0]
t = T[15]
S16 = {2: (gcb("census_timevarying"), 2016), 3: (gcb("census_timeinvariant"), 2016), 4: (gcb("baseline_trend"), 2016), 5: (gcb("census_yearwise_z"), 2016),
       7: (gmb("monitor_head_density", "beef_head"), None), 8: (gmb("monitor_farms_plus_head", "beef_cattle_farms"), None), 9: (gmb("monitor_farms_plus_head", "beef_head"), None),
       11: (gmb("monitor_farms_2019_2024", "beef_cattle_farms"), None), 12: (gmb("monitor_farms_2011_2018", "beef_cattle_farms"), None)}
for i, (r, n) in S16.items():
    N = int(n if n else r.N); setc(t, i, 2, f"{N:,}", tag="S16"); setc(t, i, 3, f"{f3(r.IRR)} ({ci(r.lo, r.hi)})", tag="S16"); setc(t, i, 4, f3(r.P_gt1), tag="S16")
# 14행(목초지 비교)은 이번 체인에 산출이 없어 그대로 둠(00_결과요약 '건너뛴 것')

# S17 (표 16): 3구간 20변수
sp = pd.read_csv(os.path.join(RS, "MASTER_final_4yr_split.csv")); ff = pd.read_csv(os.path.join(RS, "MASTER_final_4yr_fitstat.csv"))
t = T[16]; per = ["2011-2015", "2016-2019", "2020-2024"]
for j, p in enumerate(per, 1):
    n = int(ff[ff.period == p].N.iloc[0]); setc(t, 0, j, f"{p} (N = {n})", tag="S17")
for i in range(1, len(t.rows)):
    lab = t.rows[i].cells[0].text.strip().replace("1,000", "1000")
    if lab in ("Spatial fraction φ", "DIC"):
        for j, p in enumerate(per, 1):
            q = ff[ff.period == p].iloc[0]; setc(t, i, j, f"{q.phi:.3f}" if lab != "DIC" else f"{q.DIC:.1f}", tag="S17")
        continue
    for j, p in enumerate(per, 1):
        r = sp[(sp["cov"] == L2C[lab]) & (sp.period == p)].iloc[0]; setc(t, i, j, f"{f3(r.IRR)} ({ci(r.lo, r.hi)})", aa(r.cred), tag="S17")

# S18 (표 17): 사다리
ld = pd.read_csv(os.path.join(RS, "MASTER_agestd_supp_ladder.csv"))
t = T[17]
for i, r in enumerate(ld.itertuples(), 1):
    setc(t, i, 2, f"{r.DIC:.1f}", tag="S18"); setc(t, i, 3, f"{r.WAIC:.1f}", tag="S18"); setc(t, i, 4, f"{r.dDIC:.1f}", tag="S18")

# S19 (표 18): NB 대 ZINB
zb = pd.read_csv(os.path.join(RS, "MASTER_agestd_supp_zinb.csv"))
t = T[18]
for i, r in enumerate(zb.itertuples(), 1):
    setc(t, i, 1, f"{r.WAIC:.1f}", tag="S19"); setc(t, i, 2, f"{r.DIC:.1f}", tag="S19"); setc(t, i, 3, f"{r.zero_pct:.1f}", tag="S19")

# S20 (표 19): 3구간 상호작용 40개
xi = pd.read_csv(os.path.join(RS, "MASTER_final_3win_interaction.csv")).set_index("term")
lab2cov = L2C
t = T[19]
for i in range(1, len(t.rows)):
    v = lab2cov[t.rows[i].cells[0].text.strip().replace("1,000", "1000")]
    for j, w in [(1, "W2"), (2, "W3")]:
        r = xi.loc[f"{v}_x_{w}"]; setc(t, i, j, f"{f3(r.IRR)} ({ci(r.lo, r.hi)})", aa(r.cred), tag="S20")
S20_NCRED = int(xi.cred.map(tf).sum())

# S21 (표 20): 연도별 Moran, 구간 φ
mo = pd.read_csv(os.path.join(RS, "agestd_moran_yearly.csv")); pc = pd.read_csv(os.path.join(RS, "MASTER_final_covid_phi_ci.csv")).set_index("period")
lg = open(os.path.join(RR, "logs", "411b_ehsa_esri.log"), encoding="utf-8").read()
pm = re.search(r"기간평균 SIR I ([\d.]+) P ([\d.]+)", lg)
fp = lambda p: "P < 0.001" if p < 0.001 else f"P = {p:.3f}"
t = T[20]
for i, r in enumerate(mo.itertuples(), 1):
    setc(t, i, 1, f"{r.moran_I:.3f}", tag="S21"); setc(t, i, 2, f"P = {r.p:.3f}", tag="S21")
setc(t, 15, 1, f"{mo.trend_tau.iloc[0]:+.3f}", tag="S21"); setc(t, 15, 2, f"P = {mo.trend_p.iloc[0]:.3f}", tag="S21")
if pm: setc(t, 16, 1, f"{float(pm.group(1)):.3f}", tag="S21"); setc(t, 16, 2, f"P = {float(pm.group(2)):.3f}", tag="S21")
for i, (p, a, lo, hi_) in [(17, ("2011-2019", "phi_mean", "phi_lo", "phi_hi")), (18, ("2020-2024", "phi_mean", "phi_lo", "phi_hi")),
                           (19, ("2011-2019", "spatial_var_mean", "spatial_var_lo", "spatial_var_hi")), (20, ("2020-2024", "spatial_var_mean", "spatial_var_lo", "spatial_var_hi"))]:
    r = pc.loc[p]; setc(t, i, 1, f"{r[a]:.3f}", tag="S21"); setc(t, i, 2, f"{r[lo]:.3f}-{r[hi_]:.3f}", tag="S21")

# S22 (표 21): 고위험 지역
EN = {"광주시서구": ("Seo-gu", "Gwangju"), "광주시남구": ("Nam-gu", "Gwangju"), "광주시동구": ("Dong-gu", "Gwangju"), "광주시북구": ("Buk-gu", "Gwangju"), "광주시광산구": ("Gwangsan-gu", "Gwangju"),
      "전라남도나주시": ("Naju-si", "Jeollanam-do"), "인천시남구": ("Michuhol-gu", "Incheon"), "제주도제주시": ("Jeju-si", "Jeju"), "제주도서귀포시": ("Seogwipo-si", "Jeju"),
      "경상남도함안군": ("Haman-gun", "Gyeongsangnam-do"), "서울시노원구": ("Nowon-gu", "Seoul"), "서울시도봉구": ("Dobong-gu", "Seoul"), "경기도광명시": ("Gwangmyeong-si", "Gyeonggi-do"),
      "울산시동구": ("Dong-gu", "Ulsan"), "울산시남구": ("Nam-gu", "Ulsan"), "울산시중구": ("Jung-gu", "Ulsan"), "울산시북구": ("Buk-gu", "Ulsan"), "울산시울주군": ("Ulju-gun", "Ulsan"),
      "경상북도안동시": ("Andong-si", "Gyeongsangbuk-do"), "경상북도경주시": ("Gyeongju-si", "Gyeongsangbuk-do"), "전라남도함평군": ("Hampyeong-gun", "Jeollanam-do"),
      "전라남도화순군": ("Hwasun-gun", "Jeollanam-do"), "전라남도완도군": ("Wando-gun", "Jeollanam-do"), "전라남도영광군": ("Yeonggwang-gun", "Jeollanam-do"),
      "전라남도해남군": ("Haenam-gun", "Jeollanam-do"), "전라남도목포시": ("Mokpo-si", "Jeollanam-do"), "대구시달서구": ("Dalseo-gu", "Daegu"), "인천시서구": ("Seo-gu", "Incheon"),
      "충청남도홍성군": ("Hongseong-gun", "Chungcheongnam-do"), "충청남도예산군": ("Yesan-gun", "Chungcheongnam-do"), "전라남도곡성군": ("Gokseong-gun", "Jeollanam-do"),
      "전라북도순창군": ("Sunchang-gun", "Jeollabuk-do"), "경상북도봉화군": ("Bonghwa-gun", "Gyeongsangbuk-do"), "부산시동구": ("Dong-gu", "Busan")}
sr = pd.read_csv(os.path.join(RS, "MASTER_agestd_regionSIR.csv")); h = sr[sr.hi_risk == 1].sort_values("RR", ascending=False)
t = T[21]; setrows(t, 1, len(h), "S22")
for i, r in enumerate(h.itertuples(), 1):
    dn, pv = EN[r.region]; setc(t, i, 0, dn, tag="S22"); setc(t, i, 1, pv, tag="S22"); setc(t, i, 2, f3(r.RR), tag="S22"); setc(t, i, 3, ci(r.RR_lo, r.RR_hi), tag="S22")

# S23 (표 22): 지속성(영모형 100,000회, 원고 시드 20260913, v2 E) — 33 산출
ps = pd.read_csv(os.path.join(RR, "JKMS_persistence_S23_B100k_seed20260913.csv"), dtype={"sigungu_cd": str})
ps = ps[(ps.E == "E_v2") & (ps.P < 0.05)].sort_values(["P", "obs_years"], ascending=[True, False])
nm = pd.read_csv(os.path.join(os.environ.get("EHEC_GATE", os.path.join(os.environ.get("EHEC_ROOT","."),"gate")),"sigungu_key.csv"), dtype={"sigungu_cd": str}).set_index("sigungu_cd").canonical_region
t = T[22]; setrows(t, 1, len(ps), "S23")
for i, r in enumerate(ps.itertuples(), 1):
    dn, pv = EN[nm[r.sigungu_cd]]
    setc(t, i, 0, f"{dn}, {pv}", tag="S23"); setc(t, i, 1, str(int(r.obs_years)), tag="S23"); setc(t, i, 2, f"{r.exp_null:.1f}", tag="S23")
    setc(t, i, 3, "< 0.001" if r.P < 0.001 else f"{r.P:.3f}", tag="S23"); setc(t, i, 4, f"{r.FDR:.3f}", tag="S23")

# ── 다중우주(씨드 12개 완료 때만) ──
MVDONE = len(glob.glob(os.path.join(RS, "PREREG_MV_results_seed1[01][0-9].csv"))) == 12 and os.path.exists(os.path.join(RS, "PREREG_MV_summary.csv")) and \
         os.path.getmtime(os.path.join(RS, "PREREG_MV_summary.csv")) > max(os.path.getmtime(f) for f in glob.glob(os.path.join(RS, "PREREG_MV_results_seed1[01][0-9].csv")))
if MVDONE:
    MV = pd.read_csv(os.path.join(RS, "PREREG_MV_summary.csv")); MVi = MV.set_index("term")
    META = pd.read_csv(os.path.join(ROOT, "data_cache/ST_prereg_panel_META.csv")); lab_of = dict(zip(META.key, META.label))
    lab_of.update({"rum_TMid": "Ruminant density tertile, mid vs low", "rum_THigh": "Ruminant density tertile, high vs low"})
    DOMK = {"D1 reservoir": "Reservoir", "D2 agri-land": "Agriculture", "D3 water-sanit": "Water", "D4 hygiene-food": "Hygiene/food", "D5 SES": "SES", "D6 behaviour": "Behavior", "D7 detection": "Ascertainment"}
    VER = {"robust": "Robust", "robust, direction unstable": "Robust, unstable", "suggestive": "Suggestive", "null": "Null", "insufficient (n<20)": "Insufficient"}
    mv = MV[~MV.term.isin(["sex_ratio", "kchs_pc1"])].sort_values("cred_pct", ascending=False)
    t = T[12]; setrows(t, 1, len(mv), "S13")
    for i, r in enumerate(mv.itertuples(), 1):
        vals = [lab_of.get(r.term, r.term), DOMK.get(r.domain, str(r.domain)[:12]), str(int(r.n_spec)), f"{r.cred_pct:.1f}", f3(r.median_IRR), f"{r.q25:.2f}-{r.q75:.2f}", str(r.hypothesis), VER.get(r.verdict if isinstance(r.verdict, str) else ("null" if r.cred_pct < 20 else "suggestive"), str(r.verdict))]   # 2026-10-29: 534가 20% 미만 verdict를 NaN으로 남김 → 사전 규칙(<20% null)
        for j, v in enumerate(vals): setc(t, i, j, v, "a" if (j == 7 and r.cred_pct >= 50) else None, tag="S13")
    t = T[13]
    for i, k in [(6, "swine_density"), (7, "poultry_density"), (8, "fish_farm_land_area"), (9, "factory_land_area"), (10, "walking_practice")]:
        setc(t, i, 2, f"{MVi.loc[k].cred_pct:.1f}%" if k in MVi.index else "-", tag="S14")
    nspec = sum(pd.read_csv(f).spec.nunique() for f in glob.glob(os.path.join(RS, "PREREG_MV_results_seed1[01][0-9].csv")))
    S15 = pd.read_csv(os.path.join(RR, "S15_axis_robustness_v2c.csv")).set_index("term")
    t = T[14]
    for i in range(1, len(t.rows)):
        lab = t.rows[i].cells[0].text.strip(); k = [x for x in S15.index if S15.loc[x, "label"] == lab]
        if not k: continue
        r = S15.loc[k[0]]
        for j, col in enumerate(["all", "axis_no_yes", "ascertainment_none_phys_lab_all", "structure_bym2_int", "graph_q_r_knn5"], 1): setc(t, i, j, str(r[col]), tag="S15")
    LOG.append(dict(table="S12", row="", col="", before="3,000 converged (note)", after=f"수렴 사양 {nspec:,}개(각주 문장 확인)"))
else:
    LOG.append(dict(table="S12-S15", row="", col="", before="", after="다중우주 미완료 — 옛 값 그대로(씨드 12개 끝난 뒤 이 스크립트를 다시 돌릴 것)"))

# ── 새 표: 사전 지정 민감도(PC1 제외·제주 지시변수·청주 관측 0) — 번호는 저자 결정 전 임시 'S24' ──
sens = []
for sv, lab in [("nopc1", "Excluding the survey PC1 (19 covariates)"), ("jejudummy", "Adding a Jeju-si/Seogwipo-si indicator"), ("cj0", "Cheongju 2011-2013 entered as observed zeros")]:
    s2 = pd.read_csv(os.path.join(RR, f"sens_{sv}", "MASTER_final_table2.csv")).set_index("cov"); sf = dict(l.strip().split("=") for l in open(os.path.join(RR, f"sens_{sv}", "final_fit.txt")))
    hr = int(pd.read_csv(os.path.join(RR, f"sens_{sv}", "MASTER_agestd_regionSIR.csv")).hi_risk.sum())
    g = lambda k: f"{s2.loc[k].IRR:.3f} ({s2.loc[k].lo:.3f}-{s2.loc[k].hi:.3f})" + ("a" if tf(s2.loc[k].cred) else "")
    sens.append([lab, g("beef_cattle_farms"), g("fiscal_independence"), g("soap_sanitizer_use"), g("reservoir_pond_land_area"), g("current_smoking"), f"{float(sf['DIC']):.1f}", str(hr)])
mm = T2.set_index("cov"); gm = lambda k: f"{mm.loc[k].IRR:.3f} ({mm.loc[k].lo:.3f}-{mm.loc[k].hi:.3f})" + ("a" if tf(mm.loc[k].cred) else "")
sens.insert(0, ["Final model (main)", gm("beef_cattle_farms"), gm("fiscal_independence"), gm("soap_sanitizer_use"), gm("reservoir_pond_land_area"), gm("current_smoking"), f"{float(FIT['DIC']):.1f}", str(int(sr.hi_risk.sum()))])
pd.DataFrame(sens, columns=["Analysis", "Beef-cattle farm density", "Fiscal independence", "Handwashing with soap", "Reservoir and pond area", "Current smoking", "DIC", "Credibly elevated districts"]).to_csv(os.path.join(RR, "S24_sensitivity_v2c.csv"), index=False, encoding="utf-8-sig")
d.add_page_break()
p = d.add_paragraph(); r = p.add_run("Table S24 (new; number to be confirmed). Prespecified sensitivity analyses of the final 20-covariate model."); r.bold = True; r.font.size = None
hdr = ["Analysis", "Beef-cattle farm density", "Fiscal independence", "Handwashing with soap", "Reservoir and pond area", "Current smoking", "DIC", "Elevated districts, n"]
nt = d.add_table(rows=1 + len(sens), cols=len(hdr)); nt.style = d.tables[10].style
for j, hname in enumerate(hdr): nt.rows[0].cells[j].text = hname
for i, row in enumerate(sens, 1):
    for j, v in enumerate(row): nt.rows[i].cells[j].text = v
d.add_paragraph("IRR (95% CrI) per 1 SD. a, 95% credible interval excludes 1. Same model, priors, and graph as Table 2; the Jeju indicator is standardized; in the Cheongju analysis the three masked district-years enter as observed zeros.")
os.makedirs(os.path.dirname(DST), exist_ok=True); d.save(DST)
pd.DataFrame(LOG).to_csv(os.path.join(RR, "supp_table_changes_v2c.csv"), index=False, encoding="utf-8-sig")
EXTRA = dict(S7_max_hot_years=S7_MAXHOT, S7_n_at_max=S7_NMAX, S8_corr_RR_pearson=S8_CORR, S8_both_elevated=S8_BOTH, S20_ncred=S20_NCRED, MV_updated=MVDONE)
pd.Series(EXTRA).to_csv(os.path.join(RR, "supp_note_numbers_v2c.csv"), encoding="utf-8-sig")
print("saved", DST, "| changed cells", len(LOG), "|", EXTRA)
